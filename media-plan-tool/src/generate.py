#!/usr/bin/env python3
"""generate.py — メディアプラン生成スクリプト

template.xlsx の FMT タブをベースに brief.json の施策情報を反映し、
Projects/<案件名>/output/media_plan.xlsx を出力する。

使い方:
    python src/generate.py シクテン              # Projects/シクテン/input/brief.json を使用
    python src/generate.py シクテン custom.json  # 任意の JSON を指定して出力先は同じ
"""

import copy
import json
import sys
from pathlib import Path

import openpyxl

# ── パス定数 ─────────────────────────────────────────
ROOT = Path(__file__).resolve().parent.parent
TEMPLATE = ROOT / "template.xlsx"
PROJECTS = ROOT / "Projects"

# ── FMT タブ レイアウト定数 ──────────────────────────
SHEET_NAME = "メディアプラン のFMT"
FIRST_DATA = 5           # 最初のデータ行
TMPL_LAST_DATA = 15      # テンプレ最終データ行
TMPL_SUMMARY = 16        # 合計行
TMPL_MEMO = 18           # メモ行
TMPL_DATA_COUNT = TMPL_LAST_DATA - FIRST_DATA + 1  # 11

# 列番号 (1-based)
C_CAT      = 3   # C: 施策
C_DETAIL   = 4   # D: 詳細
C_ADTYPE   = 5   # E: 広告配信有無
C_TARGET   = 6   # F: ターゲット
C_SNS      = 7   # G: SNS
C_ALLOC    = 8   # H: アロケーション
C_BUDGET   = 9   # I: 予算
C_FWPRICE  = 10  # J: フォロワー単価
C_TOTALFW  = 11  # K: 合計フォロワー数
C_NUMINF   = 12  # L: 起用人数
C_AVGFW    = 13  # M: 平均フォロワー数
C_IMPRATE  = 14  # N: imp率
C_IMP      = 15  # O: imp数
C_CPM      = 16  # P: CPM
C_FQ       = 17  # Q: FQ
C_REACH    = 18  # R: リーチ数
MAX_COL    = 18

# 合計行で SUM を張る列
SUM_COLS = [
    ("H", C_ALLOC), ("K", C_TOTALFW), ("L", C_NUMINF),
    ("O", C_IMP), ("R", C_REACH),
]


# ── ユーティリティ ───────────────────────────────────
def _val(tactic, defaults, key):
    """tactic → defaults の順で値を探す。"""
    v = tactic.get(key)
    return v if v is not None else defaults.get(key)


def _save_styles(ws, row):
    """指定行の全列のスタイルを dict に退避する。"""
    out = {}
    for c in range(1, MAX_COL + 1):
        cell = ws.cell(row=row, column=c)
        out[c] = {
            "font": copy.copy(cell.font),
            "border": copy.copy(cell.border),
            "fill": copy.copy(cell.fill),
            "number_format": cell.number_format,
            "protection": copy.copy(cell.protection),
            "alignment": copy.copy(cell.alignment),
        }
    return out


def _apply_styles(ws, row, styles, height=None):
    """退避したスタイルを行に適用する。"""
    for c, s in styles.items():
        cell = ws.cell(row=row, column=c)
        for attr, val in s.items():
            setattr(cell, attr, val)
    if height is not None:
        ws.row_dimensions[row].height = height


# ── 行書込み ─────────────────────────────────────────
def _write_common(ws, r, t, summary_row):
    """全施策共通のセルを書き込む。"""
    ws.cell(row=r, column=C_CAT).value    = t["category"]
    ws.cell(row=r, column=C_DETAIL).value = t["detail"]
    ws.cell(row=r, column=C_ADTYPE).value = t["ad_type"]
    ws.cell(row=r, column=C_TARGET).value = t["target"]
    ws.cell(row=r, column=C_SNS).value    = t["sns"]
    ws.cell(row=r, column=C_ALLOC).value  = t["allocation"]
    # 予算 = 総予算 × アロケーション
    ws.cell(row=r, column=C_BUDGET).value = f"=$I${summary_row}*H{r}"


def _write_ad(ws, r, t, defaults):
    """広告配信行を書き込む。"""
    cpm = _val(t, defaults, "cpm")
    fq  = _val(t, defaults, "fq")
    if cpm is not None:
        ws.cell(row=r, column=C_CPM).value = cpm
    if fq is not None:
        ws.cell(row=r, column=C_FQ).value = fq
    # imp = 予算 / CPM × 1000
    ws.cell(row=r, column=C_IMP).value   = f"=I{r}/P{r}*1000"
    ws.cell(row=r, column=C_REACH).value = f"=O{r}/Q{r}"


def _write_influencer(ws, r, t, defaults):
    """インフルエンサー行（配信有/配信無）を書き込む。"""
    fw_price = _val(t, defaults, "follower_unit_price")
    avg_fw   = _val(t, defaults, "avg_followers")
    imp_rate = _val(t, defaults, "imp_rate")
    fq       = _val(t, defaults, "fq")
    num_inf  = t.get("num_influencers")

    # 入力値
    if fw_price is not None:
        ws.cell(row=r, column=C_FWPRICE).value = fw_price
    if avg_fw is not None:
        ws.cell(row=r, column=C_AVGFW).value = avg_fw
    if imp_rate is not None:
        ws.cell(row=r, column=C_IMPRATE).value = imp_rate
    if fq is not None:
        ws.cell(row=r, column=C_FQ).value = fq

    # 合計フォロワー数 & 起用人数
    if num_inf is not None:
        # 起用人数が明示 → 合計FW = 起用人数 × 平均FW
        ws.cell(row=r, column=C_NUMINF).value = num_inf
        if avg_fw is not None:
            ws.cell(row=r, column=C_TOTALFW).value = f"=L{r}*M{r}"
    elif fw_price is not None:
        # 標準: 合計FW = 予算 / フォロワー単価
        ws.cell(row=r, column=C_TOTALFW).value = f"=I{r}/J{r}"
        if avg_fw is not None:
            ws.cell(row=r, column=C_NUMINF).value = f"=ROUND(K{r}/M{r},0)"

    # imp = 合計FW × imp率
    ws.cell(row=r, column=C_IMP).value   = f"=K{r}*N{r}"
    # CPM（参考値）= 予算 × 1000 / imp
    ws.cell(row=r, column=C_CPM).value   = f"=I{r}*1000/O{r}"
    # リーチ = imp / FQ
    ws.cell(row=r, column=C_REACH).value = f"=O{r}/Q{r}"


# ── メイン ───────────────────────────────────────────
def generate(project_name, input_path=None):
    project_dir = PROJECTS / project_name
    if not project_dir.is_dir():
        sys.exit(f"エラー: Projects/{project_name} が見つかりません")

    path = Path(input_path) if input_path else project_dir / "input" / "brief.json"
    if not path.exists():
        sys.exit(f"エラー: {path} が見つかりません")

    with open(path, encoding="utf-8") as f:
        data = json.load(f)

    project  = data["project"]
    defaults = data.get("defaults", {})
    tactics  = data["tactics"]
    if not tactics:
        sys.exit("エラー: tactics が空です")

    total_budget = project["total_budget"]
    num          = len(tactics)

    # 出力先
    out_dir  = project_dir / "output"
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "media_plan.xlsx"

    # テンプレート読込
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[SHEET_NAME]

    # スタイル退避（行挿入/削除の前に）
    base_styles = _save_styles(ws, FIRST_DATA)
    row_height  = ws.row_dimensions[FIRST_DATA].height

    # ── 行数調整 ──
    diff = num - TMPL_DATA_COUNT
    if diff > 0:
        # 合計行の直前に行を挿入
        ws.insert_rows(TMPL_LAST_DATA + 1, diff)
    elif diff < 0:
        # 不要な行を削除（データ末尾側から）
        ws.delete_rows(FIRST_DATA + num, -diff)

    summary_row   = TMPL_SUMMARY + diff
    memo_row      = TMPL_MEMO + diff
    last_data_row = FIRST_DATA + num - 1

    # ── 総予算セット ──
    ws.cell(row=summary_row, column=C_BUDGET).value = total_budget

    # ── 施策の書込み ──
    for i, t in enumerate(tactics):
        r = FIRST_DATA + i
        is_ad = t["ad_type"] == "広告配信"

        # スタイル適用 → セルクリア → 値書込み
        _apply_styles(ws, r, base_styles, row_height)
        for col in range(C_CAT, MAX_COL + 1):
            ws.cell(row=r, column=col).value = None

        _write_common(ws, r, t, summary_row)
        if is_ad:
            _write_ad(ws, r, t, defaults)
        else:
            _write_influencer(ws, r, t, defaults)

    # ── 合計行の SUM 数式を更新 ──
    for letter, col in SUM_COLS:
        ws.cell(row=summary_row, column=col).value = (
            f"=SUM({letter}{FIRST_DATA}:{letter}{last_data_row})"
        )

    # ── メモ ──
    memo = project.get("memo")
    if memo:
        ws.cell(row=memo_row, column=C_DETAIL).value = memo

    # ── 保存 ──
    wb.save(out_path)
    print(f"生成完了: {out_path}")


if __name__ == "__main__":
    if len(sys.argv) < 2:
        sys.exit("使い方: python src/generate.py <project_name> [input.json]")
    pname = sys.argv[1]
    ipath = sys.argv[2] if len(sys.argv) > 2 else None
    generate(pname, ipath)
